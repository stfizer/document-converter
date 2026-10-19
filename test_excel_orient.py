# -*- coding: utf-8 -*-
"""Excel→PDF 页面方向/列缩放验证"""
import os
import sys

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE)
TD = os.path.join(BASE, "test_data", "orient")
os.makedirs(TD, exist_ok=True)

import docflow.converters as cv
import pymupdf

HEADERS = ["序号", "部门代码", "用人单位", "职位代码", "考试类别", "招考人数",
           "专业（本科及以上）", "面试人员比例", "工作地点", "备注"]
ROWS = [
    ["7", "130121", "国家税务总局武汉市青山区税务局", "30011013002", "行政执法类", "3",
     "经济学类、财政学类、金融学类、会计审计、社会保障", "3:1", "湖北省武汉市青山区",
     "2027届高校毕业生；本单位不提供宿舍，在本单位最低服务年限为5年"],
    ["9", "130121", "国家税务总局武汉市洪山区税务局", "30011014001", "行政执法类", "2",
     "经济学类、财政学类、金融学类、会计审计", "3:1", "湖北省武汉市洪山区",
     "2027届高校毕业生；在本单位最低服务年限为5年"],
    ["11", "130121", "国家税务总局武汉市市税务局", "30011015004", "行政执法类", "4",
     "经济学类、财政学类、金融学类", "3:1", "湖北省武汉市市辖区",
     "2027届高校毕业生；服务年限为5年"],
    ["15", "130121", "国家税务总局武汉市江岸区税务局", "30011016001", "行政执法类", "3",
     "经济学类、经济与贸易类、社会保障", "3:1", "湖北省武汉市江岸区", "最低服务年限5年"],
    ["17", "130121", "国家税务总局武汉市江汉区税务局", "30011017002", "行政执法类", "2",
     "经济学类、经济与贸易类", "3:1", "湖北省武汉市江汉区", "服务年限5年"],
]


def make_wide_xlsx():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    p = os.path.join(TD, "wide_table.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "职位表"
    ws.append(HEADERS)
    for row in ROWS:
        ws.append(row)
    hf = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="2563EB")
    for c in ws[1]:
        c.font = hf
        c.fill = hfill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [6, 10, 34, 14, 11, 9, 42, 11, 18, 38]
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    wb.save(p)
    return p


def check(name, cond, extra=""):
    print(f"[{'PASS' if cond else 'FAIL'}] {name} {extra}")
    return cond


def page_info(pdf_path):
    d = pymupdf.open(pdf_path)
    p = d[0]
    w, h = p.rect.width, p.rect.height     # 必须在 close 前取值
    t = p.get_text()
    d.close()
    return w, h, t


def main():
    src = make_wide_xlsx()
    import time
    time.sleep(1.5)   # 等 Excel COM 能稳定打开刚保存的文件
    log = lambda m: print("   ", m)
    all_ok = True

    # a) 跟随文档设置 (Excel 默认纵向, 列会被拆页)
    d1 = os.path.join(TD, "out_auto.pdf")
    ok, msg = cv.excel_to_pdf(src, d1, log=log)
    print("AUTO ok=", ok, "msg=", msg,
          "size=", os.path.getsize(d1) if os.path.exists(d1) else "MISSING")
    w, h, t = page_info(d1)
    print(f"   默认: {w:.0f}x{h:.0f}pt 纵向={h > w}")
    all_ok &= check("默认模式输出 PDF", os.path.exists(d1) and os.path.getsize(d1) > 0)

    # b) 横向 + 列缩放一页宽
    d2 = os.path.join(TD, "out_landscape.pdf")
    cv.excel_to_pdf(src, d2, log=log, landscape=True, fit_wide=True)
    w, h, t = page_info(d2)
    has_last_col = ("备注" in t) and ("2027届高校毕业生" in t or "服务年限" in t)
    print(f"   横向: {w:.0f}x{h:.0f}pt 横向={w > h} 首页含最右列内容={has_last_col}")
    all_ok &= check("横向页面 (宽>高)", w > h, f"{w:.0f}x{h:.0f}")
    all_ok &= check("最右列(备注)在第1页", has_last_col)

    # c) 纵向 + 列缩放一页宽 (列完整但方向不变)
    d3 = os.path.join(TD, "out_portrait_fit.pdf")
    cv.excel_to_pdf(src, d3, log=log, landscape=False, fit_wide=True)
    w, h, t = page_info(d3)
    print(f"   纵向+缩放: {w:.0f}x{h:.0f}pt 纵向={h > w} 含最右列={'备注' in t}")
    all_ok &= check("纵向保持 (高>宽)", h > w, f"{w:.0f}x{h:.0f}")
    all_ok &= check("纵向+缩放后列完整", "备注" in t)

    # d) 兜底引擎横向 (无 Office 场景)
    d4 = os.path.join(TD, "out_fallback.pdf")
    cv._xlsx_text_fallback_pdf(src, d4, landscape=True)
    w, h, t = page_info(d4)
    print(f"   兜底横向: {w:.0f}x{h:.0f}pt 横向={w > h}")
    all_ok &= check("兜底引擎横向页面", w > h, f"{w:.0f}x{h:.0f}")

    print("\n=====", "全部通过" if all_ok else "存在失败", "=====")
    return 0 if all_ok else 1


if __name__ == "__main__":
    sys.exit(main())

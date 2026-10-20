# -*- coding: utf-8 -*-
"""Excel→PDF 总页数均分功能验证"""
import os
import sys

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE)
os.chdir(BASE)

import docflow.converters as cv
import pymupdf

TD = os.path.join(BASE, "test_data", "orient")


def make_long_xlsx():
    """60 行 × 10 列宽表 (自然打印约 3-4 页纵向内容)"""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    p = os.path.join(TD, "long_table.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["序号", "部门代码", "用人单位", "职位代码", "考试类别",
               "招考人数", "专业", "面试比例", "工作地点", "备注"])
    for r in range(1, 61):
        ws.append([str(r), f"1301{r:02d}", f"测试单位第{r}号（名称较长以撑宽表格）",
                   f"300110{r:05d}", "行政执法类", str(r % 5 + 1),
                   f"经济学类、财政学类、金融学类专业{r}", "3:1",
                   f"湖北省武汉市某区{r}", f"备注内容{r}，服务年限5年，2027届毕业生"])
    hf = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="2563EB")
    for c in ws[1]:
        c.font = hf
        c.fill = hfill
    widths = [6, 10, 30, 14, 11, 9, 34, 10, 16, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(p)
    return p


def pages(path):
    d = pymupdf.open(path)
    texts = [d[i].get_text() for i in range(len(d))]
    w, h = d[0].rect.width, d[0].rect.height
    d.close()
    return len(texts), texts, w, h


def check(name, cond, extra=""):
    print(f"[{'PASS' if cond else 'FAIL'}] {name} {extra}")
    return cond


def main():
    src = make_long_xlsx()
    import time
    time.sleep(1.2)
    log = lambda m: print("   ", m)
    ok_all = True

    # a) 横向 + 列缩放 + 总页数 3
    d1 = os.path.join(TD, "tp3.pdf")
    cv.excel_to_pdf(src, d1, log=log, landscape=True, fit_wide=True, total_pages=3)
    n, texts, w, h = pages(d1)
    print(f"   横向+3页: {n} 页, {w:.0f}x{h:.0f}")
    ok_all &= check("输出恰好 3 页", n == 3, f"实际 {n} 页")
    ok_all &= check("横向页面", w > h, f"{w:.0f}x{h:.0f}")
    ok_all &= check("第 1 页含首部内容", "测试单位第1号" in texts[0])
    ok_all &= check("末页含尾部内容", "测试单位第60号" in texts[-1])
    ok_all &= check("首页含最右列(备注)", "备注内容1" in texts[0])

    # b) 横向 + 列缩放 + 总页数 1 (整表一页)
    d2 = os.path.join(TD, "tp1.pdf")
    cv.excel_to_pdf(src, d2, log=log, landscape=True, fit_wide=True, total_pages=1)
    n, texts, w, h = pages(d2)
    print(f"   横向+1页: {n} 页")
    ok_all &= check("输出恰好 1 页", n == 1, f"实际 {n} 页")
    ok_all &= check("整表单页含首尾内容",
                    "测试单位第1号" in texts[0] and "测试单位第60号" in texts[0])

    # c) 自动分页 (total_pages=0/None) 回归
    d3 = os.path.join(TD, "tp_auto.pdf")
    cv.excel_to_pdf(src, d3, log=log, landscape=True, fit_wide=True, total_pages=None)
    n, texts, w, h = pages(d3)
    print(f"   自动分页: {n} 页")
    ok_all &= check("自动模式正常输出", n >= 2, f"实际 {n} 页")

    # d) 兜底引擎按行均分 (行数均分, 页数可能因内容高度多于 N)
    d4 = os.path.join(TD, "tp_fallback.pdf")
    cv._xlsx_text_fallback_pdf(src, d4, landscape=True, total_pages=3)
    n, texts, w, h = pages(d4)
    all_text = "".join(texts)
    print(f"   兜底均分: {n} 页, {w:.0f}x{h:.0f}")
    ok_all &= check("兜底横向输出", n >= 3 and w > h, f"实际 {n} 页")
    ok_all &= check("兜底含首部内容", "测试单位第1号" in all_text)
    ok_all &= check("兜底含尾部内容", "测试单位第60号" in all_text)

    print("\n=====", "全部通过" if ok_all else "存在失败", "=====")
    return 0 if ok_all else 1


if __name__ == "__main__":
    sys.exit(main())

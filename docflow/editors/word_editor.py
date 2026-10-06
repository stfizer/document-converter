# -*- coding: utf-8 -*-
"""Word (.docx) 文字与图片编辑页 - Qt 版"""
import os

from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QListWidget, QPushButton,
    QFileDialog, QMessageBox, QTextEdit, QSplitter, QGroupBox,
)
from PySide6.QtCore import Qt

from docflow.drops import enable_page_drop


class WordEditorPage(QWidget):
    def __init__(self, status_cb, parent=None):
        super().__init__(parent)
        self.status_cb = status_cb
        self.doc = None
        self.path = None
        self._build()
        enable_page_drop(self, self._drop_open, (".docx",))

    def _drop_open(self, files):
        ok = skipped = 0
        for p in files:
            if p.lower().endswith(".docx"):
                self.open_path(p)
                ok += 1
            else:
                skipped += 1
        if skipped:
            self.status_cb(f"忽略 {skipped} 个非 .docx 文件（旧版 .doc 请先另存为 .docx）")

    def _build(self):
        lay = QVBoxLayout(self)
        lay.setContentsMargins(16, 12, 16, 12)
        lay.setSpacing(10)

        bar = QHBoxLayout()
        b_open = QPushButton("📂 打开 Word")
        b_open.clicked.connect(self.open_file)
        b_save = QPushButton("💾 另存为")
        b_save.clicked.connect(self.save_as)
        self.lbl_file = QLabel("未打开文件")
        self.lbl_file.setStyleSheet("color:#64748b;")
        bar.addWidget(b_open)
        bar.addWidget(b_save)
        bar.addWidget(self.lbl_file, stretch=1)
        lay.addLayout(bar)

        g1 = QGroupBox("段落列表（点击选中）")
        v1 = QVBoxLayout(g1)
        self.lst = QListWidget()
        self.lst.setAlternatingRowColors(True)
        self.lst.currentRowChanged.connect(self._on_select)
        v1.addWidget(self.lst)
        lay.addWidget(g1, stretch=3)

        g2 = QGroupBox("编辑选中段落")
        v2 = QVBoxLayout(g2)
        self.txt = QTextEdit()
        self.txt.setPlaceholderText("←  先打开 Word 文档，再从上方选择段落进行编辑…")
        self.txt.setMinimumHeight(120)
        v2.addWidget(self.txt)
        h = QHBoxLayout()
        for text, cb in (("✔ 应用修改到段落", self._apply_text),
                         ("➕ 下方新增段落", self._add_para),
                         ("🖼 选中段落末插入图片", self._insert_image),
                         ("🗑 删除段落", self._del_para)):
            b = QPushButton(text)
            b.clicked.connect(cb)
            h.addWidget(b)
        h.addStretch(1)
        self.lbl_msg = QLabel("")
        self.lbl_msg.setStyleSheet("color:#2563eb;")
        h.addWidget(self.lbl_msg, stretch=1)
        v2.addLayout(h)
        lay.addWidget(g2, stretch=2)

    # ------------------------------------------------ 数据
    def open_file(self):
        p, _ = QFileDialog.getOpenFileName(self, "打开 Word 文档", "",
                                           "Word 文档 (*.docx);;所有文件 (*.*)")
        if p:
            self.open_path(p)

    def open_path(self, p):
        try:
            from docx import Document
            self.doc = Document(p)
        except Exception as e:
            QMessageBox.critical(self, "打开失败", f"无法打开该文件:\n{e}")
            return
        self.path = p
        self.lbl_file.setText(p)
        self._reload()
        self.status_cb(f"已打开 {p}")

    def _reload(self):
        self.lst.blockSignals(True)
        self.lst.clear()
        for i, p in enumerate(self.doc.paragraphs):
            t = p.text.strip()
            mark = "[空行]" if not t else (t[:52] + ("…" if len(t) > 52 else ""))
            self.lst.addItem(f"{i:>3}.  {mark}")
        self.lst.blockSignals(False)

    def _selected(self):
        row = self.lst.currentRow()
        if row < 0:
            QMessageBox.information(self, "提示", "请先在列表中选择一个段落")
            return None
        return row

    def _on_select(self, row):
        if row < 0 or self.doc is None:
            return
        self.txt.setPlainText(self.doc.paragraphs[row].text)

    # ------------------------------------------------ 操作
    def _require_doc(self):
        if self.doc is None:
            QMessageBox.information(self, "提示", "请先打开一个 Word 文档")
            return False
        return True

    def _apply_text(self):
        if not self._require_doc():
            return
        idx = self._selected()
        if idx is None:
            return
        new = self.txt.toPlainText()
        p = self.doc.paragraphs[idx]
        try:
            if p.runs:
                p.runs[0].text = new
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.add_run(new)
        except Exception as e:
            QMessageBox.critical(self, "错误", f"修改失败:\n{e}")
            return
        self._reload()
        self.lst.setCurrentRow(idx)
        self.lbl_msg.setText(f"已修改第 {idx} 段（保存后生效）")
        self.status_cb(f"已修改第 {idx} 段（保存后生效）")

    def _add_para(self):
        if not self._require_doc():
            return
        idx = self._selected()
        if idx is None:
            return
        text = self.txt.toPlainText()
        try:
            from docx.oxml.ns import qn
            ref = self.doc.paragraphs[idx]
            new_p = ref._p.makeelement(qn("w:p"), {})
            ref._p.addnext(new_p)
            from docx.text.paragraph import Paragraph
            para = Paragraph(new_p, ref._parent)
            para.add_run(text)
        except Exception as e:
            QMessageBox.critical(self, "错误", f"插入失败:\n{e}")
            return
        self._reload()
        self.lst.setCurrentRow(idx + 1)
        self.status_cb(f"已在第 {idx} 段后新增段落")

    def _insert_image(self):
        if not self._require_doc():
            return
        idx = self._selected()
        if idx is None:
            return
        p, _ = QFileDialog.getOpenFileName(self, "选择图片", "",
                                           "图片 (*.png *.jpg *.jpeg *.bmp *.gif *.webp);;所有文件 (*.*)")
        if not p:
            return
        try:
            from docx.shared import Inches
            para = self.doc.paragraphs[idx]
            para.add_run().add_picture(p, width=Inches(4.0))
        except Exception as e:
            QMessageBox.critical(self, "错误", f"插入图片失败:\n{e}")
            return
        self._reload()
        self.lst.setCurrentRow(idx)
        self.lbl_msg.setText(f"已在第 {idx} 段末尾插入图片: {os.path.basename(p)}")
        self.status_cb("图片已插入（保存后生效）")

    def _del_para(self):
        if not self._require_doc():
            return
        idx = self._selected()
        if idx is None:
            return
        if QMessageBox.question(self, "确认", f"确定删除第 {idx} 段吗？") != QMessageBox.Yes:
            return
        p = self.doc.paragraphs[idx]
        p._element.getparent().remove(p._element)
        self._reload()
        self.status_cb(f"已删除第 {idx} 段")

    def save_as(self):
        if not self._require_doc():
            return
        p, _ = QFileDialog.getSaveFileName(self, "另存 Word 文档", os.path.basename(self.path),
                                           "Word 文档 (*.docx)")
        if not p:
            return
        try:
            self.doc.save(p)
        except PermissionError:
            QMessageBox.critical(self, "保存失败", "文件被占用（可能正被 Word 打开），请先关闭后重试。")
            return
        self.status_cb(f"已保存: {p}")
        QMessageBox.information(self, "保存成功", f"已保存到:\n{p}")

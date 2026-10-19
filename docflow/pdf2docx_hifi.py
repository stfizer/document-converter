# -*- coding: utf-8 -*-
"""PDF -> Word 高保真转换模式

原理: 用 PyMuPDF 提取每个文本行/图片/矢量图形的精确位置、字体、字号、颜色,
在 Word 中用"页面绝对定位的文本框/形状"逐元素重建, 视觉上与原 PDF 几乎一致。
- 下划线: PDF 中的下划线是独立画线, 自动关联到其上方文字行, 转为 Word 真实下划线格式
- 加粗: 字体 flags + 字体名启发式 (Bold/Heavy/Black)
- 中文字体名映射: SimSun->宋体 等, 提高Word字体命中率
适合存档/打印/查看; 局限: 文字位于文本框中, 大段改写不如普通文档方便。
"""
import io
import re

import pymupdf
from docx import Document
from docx.shared import Emu
from docx.oxml import parse_xml

EMU = 12700  # 1pt = 12700 EMU

_BOLD_RE = re.compile(r"(bold|heavy|black|extrabold|ultrabold|semibold|demibold)", re.I)

FONT_MAP = {
    "SimSun": "宋体", "NSimSun": "新宋体", "SimHei": "黑体",
    "KaiTi": "楷体", "KaiTi_GB2312": "楷体_GB2312",
    "FangSong": "仿宋", "FangSong_GB2312": "仿宋_GB2312",
    "MicrosoftYaHei": "微软雅黑", "Microsoft YaHei": "微软雅黑",
    "DengXian": "等线", "STZhongsong": "华文中宋", "STKaiti": "华文楷体",
    "STXihei": "华文细黑", "STSong": "华文宋体", "STFangsong": "华文仿宋",
    "LiSu": "隶书", "YouYuan": "幼圆",
}


def _map_font(name):
    return FONT_MAP.get(name, FONT_MAP.get(name.replace(" ", ""), name))

NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
)

_nid_counter = [100]


def _nid():
    _nid_counter[0] += 1
    return _nid_counter[0]


def _esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _hex(c):
    return "%06X" % (int(c) & 0xFFFFFF)


def _hexf(seq):
    return _hex((int(seq[0] * 255) << 16) | (int(seq[1] * 255) << 8) | int(seq[2] * 255))


def _clean_font(name, flags):
    """去掉子集前缀(如 ABCDEF+SimSun); 提取粗斜体 (flags + 字体名启发式)"""
    name = name.split("+")[-1]
    bold = bool(flags & 16)
    italic = bool(flags & 2)
    for suf, b, i in (("-BoldItalic", 1, 1), ("-BoldMT", 1, 0), ("-Bold", 1, 0),
                      ("-ItalicMT", 0, 1), ("-Oblique", 0, 1), ("-Italic", 0, 1)):
        if name.endswith(suf):
            name = name[: -len(suf)]
            bold = bold or bool(b)
            italic = italic or bool(i)
            break
    if _BOLD_RE.search(name):
        bold = True
    return _map_font(name), bold, italic


def _line_textbox(line, ul_rects=()):
    """一个文本行 -> 一个页面绝对定位的无边框文本框
    ul_rects: 命中该行的下划线线段 x 区间列表, 命中的 span 加真实下划线格式"""
    x0, y0, x1, y1 = line["bbox"]
    spans = [s for s in line.get("spans", []) if s.get("text", "").strip()]
    if not spans:
        return None, []
    size = max(s.get("size", 11) for s in spans)
    w = (x1 - x0) + size * 0.2 + 1
    h = (y1 - y0) + size * 0.25 + 1
    line_tw = max(20, int(round((y1 - y0) * 20)))

    runs = []
    used_ul = []
    for sp in spans:
        sb = sp.get("bbox", (x0, y0, x1, y1))
        sw_ = max(sb[2] - sb[0], 0.1)
        ul = any(min(sb[2], u[1]) - max(sb[0], u[0]) > 0.35 * sw_ for u in ul_rects)
        if ul:
            used_ul.append((sb[0], sb[2]))
        font, bold, italic = _clean_font(sp.get("font", "Arial"), sp.get("flags", 0))
        color = _hex(sp.get("color", 0))
        sz = max(2, int(round(sp.get("size", 11) * 2)))
        fmt = ("<w:b/>" if bold else "") + ("<w:i/>" if italic else "") + \
              ('<w:u w:val="single"/>' if ul else "")
        runs.append(
            f'<w:r><w:rPr><w:rFonts w:ascii="{font}" w:hAnsi="{font}" '
            f'w:eastAsia="{font}"/>{fmt}<w:color w:val="{color}"/>'
            f'<w:sz w:val="{sz}"/><w:szCs w:val="{sz}"/></w:rPr>'
            f'<w:t xml:space="preserve">{_esc(sp["text"])}</w:t></w:r>')
    para = (f'<w:p><w:pPr><w:spacing w:before="0" w:after="0" '
            f'w:line="{line_tw}" w:lineRule="exact"/></w:pPr>{"".join(runs)}</w:p>')

    shape = "<a:noFill/><a:ln><a:noFill/></a:ln>"
    z = 1000000 + _nid()
    e = lambda v: int(round(v * EMU))
    nid = _nid()
    xml = (f'<w:r {NS}><w:drawing>'
           f'<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" '
           f'relativeHeight="{z}" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
           f'<wp:simplePos x="0" y="0"/>'
           f'<wp:positionH relativeFrom="page"><wp:posOffset>{e(x0)}</wp:posOffset></wp:positionH>'
           f'<wp:positionV relativeFrom="page"><wp:posOffset>{e(y0)}</wp:posOffset></wp:positionV>'
           f'<wp:extent cx="{e(w)}" cy="{e(h)}"/>'
           f'<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
           f'<wp:docPr id="{nid}" name="tb{nid}"/><wp:cNvGraphicFramePr/>'
           f'<a:graphic><a:graphicData '
           f'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
           f'<wps:wsp><wps:cNvSpPr txBox="1"/>'
           f'<wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{e(w)}" cy="{e(h)}"/></a:xfrm>'
           f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>{shape}</wps:spPr>'
           f'<wps:txbx><w:txbxContent>{para}</w:txbxContent></wps:txbx>'
           f'<wps:bodyPr rot="0" wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" '
           f'anchor="t" anchorCtr="0"><a:noAutofit/></wps:bodyPr>'
           f'</wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>')
    return xml, used_ul


def _rect_shape(x, y, w, h, fill_hex=None, stroke_hex=None, stroke_w=1.0):
    """填充矩形/线条 -> Word 矩形形状 (置于文字与图片下层)"""
    if fill_hex:
        fill = f'<a:solidFill><a:srgbClr val="{fill_hex}"/></a:solidFill>'
    else:
        fill = "<a:noFill/>"
    if stroke_hex:
        ln = (f'<a:ln w="{max(1270, int(stroke_w * EMU))}">'
              f'<a:solidFill><a:srgbClr val="{stroke_hex}"/></a:solidFill></a:ln>')
    else:
        ln = "<a:ln><a:noFill/></a:ln>"
    z = _nid()
    e = lambda v: int(round(v * EMU))
    nid = _nid()
    return (f'<w:r {NS}><w:drawing>'
            f'<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" '
            f'relativeHeight="{z}" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">'
            f'<wp:simplePos x="0" y="0"/>'
            f'<wp:positionH relativeFrom="page"><wp:posOffset>{e(x)}</wp:posOffset></wp:positionH>'
            f'<wp:positionV relativeFrom="page"><wp:posOffset>{e(y)}</wp:posOffset></wp:positionV>'
            f'<wp:extent cx="{e(max(w, 0.3))}" cy="{e(max(h, 0.3))}"/>'
            f'<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
            f'<wp:docPr id="{nid}" name="rc{nid}"/><wp:cNvGraphicFramePr/>'
            f'<a:graphic><a:graphicData '
            f'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
            f'<wps:wsp><wps:cNvSpPr/>'
            f'<wps:spPr><a:xfrm><a:off x="0" y="0"/>'
            f'<a:ext cx="{e(max(w, 0.3))}" cy="{e(max(h, 0.3))}"/></a:xfrm>'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>{fill}{ln}</wps:spPr>'
            f'<wps:bodyPr rot="0" anchor="t" anchorCtr="0"/>'
            f'</wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>')


def _pic_run(rid, x, y, w, h):
    """图片 -> 页面绝对定位的浮动图片 (文字下层)"""
    e = lambda v: int(round(v * EMU))
    nid = _nid()
    return (f'<w:r {NS}><w:drawing>'
            f'<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" '
            f'relativeHeight="{500000 + nid}" behindDoc="1" locked="0" '
            f'layoutInCell="1" allowOverlap="1">'
            f'<wp:simplePos x="0" y="0"/>'
            f'<wp:positionH relativeFrom="page"><wp:posOffset>{e(x)}</wp:posOffset></wp:positionH>'
            f'<wp:positionV relativeFrom="page"><wp:posOffset>{e(y)}</wp:posOffset></wp:positionV>'
            f'<wp:extent cx="{e(w)}" cy="{e(h)}"/>'
            f'<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
            f'<wp:docPr id="{nid}" name="img{nid}"/><wp:cNvGraphicFramePr/>'
            f'<a:graphic><a:graphicData '
            f'uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
            f'<pic:pic><pic:nvPicPr><pic:cNvPr id="{nid}" name="img{nid}"/>'
            f'<pic:cNvPicPr/></pic:nvPicPr>'
            f'<pic:blipFill><a:blip r:embed="{rid}"/>'
            f'<a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
            f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{e(w)}" cy="{e(h)}"/></a:xfrm>'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>'
            f'</pic:pic></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>')


def pdf_to_docx_hifi(src, dst, log=print):
    pdf = pymupdf.open(src)
    docx = Document()
    for p0 in list(docx.paragraphs):        # 清空默认空段
        p0._element.getparent().remove(p0._element)

    n_pages = n_tb = n_pic = n_shape = 0
    for pno, page in enumerate(pdf):
        sec = docx.sections[0] if pno == 0 else docx.add_section()
        sec.page_width = Emu(int(round(page.rect.width * EMU)))
        sec.page_height = Emu(int(round(page.rect.height * EMU)))
        sec.left_margin = sec.right_margin = Emu(0)
        sec.top_margin = sec.bottom_margin = Emu(0)
        sec.header_distance = sec.footer_distance = Emu(0)

        runs = []
        # ---- 矢量图形分类: 水平细线作为"下划线候选", 其余待画形状
        try:
            drawings = page.get_drawings()
        except Exception:
            drawings = []
        thin = []  # (idx, rect, color) 水平细描边线/细填充条: 下划线候选(或表格横线)
        for di, d in enumerate(drawings):
            r = d.get("rect")
            if r is None:
                continue
            # 注意: 线条的包围盒宽或高为 0 (Rect.is_empty=True), 是有效元素, 不能过滤
            if r.width > page.rect.width + 1 or r.height > page.rect.height + 1:
                continue
            fill, stroke = d.get("fill"), d.get("color")
            sw = d.get("width") or 1.0
            try:
                if stroke is not None and "s" in d["type"] and "f" not in d["type"] \
                        and r.height <= max(2.5, sw * 2) \
                        and 4 <= r.width <= page.rect.width - 5:
                    thin.append((di, r, stroke))
                elif fill is not None and "f" in d["type"] and "s" not in d["type"] \
                        and r.height <= 2.2 and r.width >= 4:
                    thin.append((di, r, fill))
            except Exception:
                continue

        # ---- 图片
        try:
            infos = page.get_image_info(xrefs=True)
        except Exception:
            infos = []
        for info in infos:
            if not info.get("xref"):
                continue
            try:
                raw = pdf.extract_image(info["xref"])
                rid, _img = docx.part.get_or_add_image(io.BytesIO(raw["image"]))
                bx = info["bbox"]
                runs.append(_pic_run(rid, bx[0], bx[1], bx[2] - bx[0], bx[3] - bx[1]))
                n_pic += 1
            except Exception as e:
                log(f"  ⚠ 第{pno + 1}页图片提取失败: {e}")

        # ---- 文本行 (最上层), 并把紧贴行底的水平线关联为 Word 真实下划线
        ul_used = set()
        for block in page.get_text("dict")["blocks"]:
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                x0, y0, x1, y1 = line["bbox"]
                ul_rects = []
                for di, r, color in thin:
                    if di in ul_used:
                        continue
                    if y1 - 1.0 <= r.y0 <= y1 + 4.5:      # 线紧贴该行底部
                        overlap = min(x1, r.x1) - max(x0, r.x0)
                        if overlap >= 0.5 * max(1e-6, min(x1 - x0, r.x1 - r.x0)):
                            ul_rects.append((r.x0, r.x1))
                            ul_used.add(di)
                tb, _used = _line_textbox(line, ul_rects)
                if tb:
                    runs.append(tb)
                    n_tb += 1

        # ---- 剩余图形 (未匹配为下划线的表格横线/色块/外框等), 置于底层
        for di, d in enumerate(drawings):
            if di in ul_used:
                continue
            r = d.get("rect")
            if r is None:
                continue
            if r.width > page.rect.width + 1 or r.height > page.rect.height + 1:
                continue
            fill, stroke = d.get("fill"), d.get("color")
            sw = d.get("width") or 1.0
            try:
                if fill is not None and "f" in d["type"]:
                    runs.append(_rect_shape(r.x0, r.y0, r.width, r.height,
                                            fill_hex=_hexf(fill)))
                    n_shape += 1
                elif stroke is not None and "s" in d["type"]:
                    if r.height <= max(2.5, sw * 2):     # 水平细线
                        runs.append(_rect_shape(r.x0, r.y0, r.width,
                                                max(r.height, sw),
                                                fill_hex=_hexf(stroke)))
                    else:
                        runs.append(_rect_shape(r.x0, r.y0, r.width, r.height,
                                                stroke_hex=_hexf(stroke), stroke_w=sw))
                    n_shape += 1
            except Exception:
                continue

        # ---- 本页所有元素挂在一个锚定空段落上
        p = docx.add_paragraph()
        p._p.append(parse_xml(
            '<w:pPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:spacing w:before="0" w:after="0" w:line="14" w:lineRule="exact"/>'
            '<w:rPr><w:sz w:val="2"/></w:rPr></w:pPr>'))
        for rx in runs:
            try:
                p._p.append(parse_xml(rx))
            except Exception as e:
                log(f"  ⚠ 元素解析跳过: {e}")
        n_pages += 1

    pdf.close()
    docx.save(dst)
    return True, (f"高保真转换完成: {n_pages} 页, {n_tb} 文本行, "
                  f"{n_pic} 图片, {n_shape} 图形 (文字位于文本框中)")

# -*- coding: utf-8 -*-
"""Excel (.xlsx) 单元格文字与图片编辑页 - Qt 版"""
import os

from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton, QFileDialog,
    QMessageBox, QTextEdit, QComboBox, QTableWidget, QTableWidgetItem,
    QGroupBox, QHeaderView, QAbstractItemView,
)
from PySide6.QtCore import Qt

from openpyxl.utils import get_column_letter

from docflow.drops import enable_page_drop


class ExcelEditorPage(QWidget):
    PREVIEW_ROWS = 500
    PREVIEW_COLS = 60

    def __init__(self, status_cb, parent=None):
        super().__init__(parent)
        self.status_cb = status_cb
        self.wb = None
        self.ws = None
        self.path = None
        self._img_count = 0
        self._build()
        enable_page_drop(self, self._drop_open, (".xlsx", ".xlsm"))

    def _drop_open(self, files):
        ok = skipped = 0
        for p in files:
            if p.lower().endswith((".xlsx", ".xlsm")):
                self.open_path(p)
                ok += 1
            else:
                skipped += 1
        if skipped:
            self.status_cb(f"忽略 {skipped} 个非 Excel 文件（旧版 .xls 请先另存为 .xlsx）")

    def _build(self):
        lay = QVBoxLayout(self)
        lay.setContentsMargins(16, 12, 16, 12)
        lay.setSpacing(10)

        bar = QHBoxLayout()
        b_open = QPushButton("📂 打开 Excel")
        b_open.clicked.connect(self.open_file)
        b_save = QPushButton("💾 另存为")
        b_save.clicked.connect(self.save_as)
        self.lbl_file = QLabel("未打开文件")
        self.lbl_file.setStyleSheet("color:#64748b;")
        bar.addWidget(b_open)
        bar.addWidget(b_save)
        self.lbl_file = QLabel("未打开文件")
        self.lbl_file.setStyleSheet("color:#64748b;")
        bar.addWidget(self.lbl_file, stretch=1)
        bar.addWidget(QLabel("工作表:"))
        self.cmb_sheet = QComboBox()
        self.cmb_sheet.setMinimumWidth(160)
        self.cmb_sheet.currentTextChanged.connect(self._show_sheet)
        bar.addWidget(self.cmb_sheet)
        lay.addLayout(bar)

        split = QHBoxLayout()
        split.setSpacing(10)

        g1 = QGroupBox("工作表内容（点击单元格选中）")
        v1 = QVBoxLayout(g1)
        self.table = QTableWidget()
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionBehavior(QAbstractItemView.SelectItems)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.currentCellChanged.connect(self._on_cell)
        v1.addWidget(self.table)
        split.addWidget(g1, stretch=3)

        g2 = QGroupBox("编辑选中单元格")
        v2 = QVBoxLayout(g2)
        self.lbl_pos = QLabel("单元格: -")
        self.lbl_pos.setStyleSheet("font-weight:bold;")
        v2.addWidget(self.lbl_pos)
        self.txt = QTextEdit()
        self.txt.setMaximumHeight(110)
        v2.addWidget(self.txt)
        b_apply = QPushButton("✔ 应用修改")
        b_apply.clicked.connect(self._apply_cell)
        v2.addWidget(b_apply)

        sep = QLabel(" ")
        v2.addWidget(sep)
        self.lbl_anchor = QLabel("图片锚点: -")
        self.lbl_anchor.setStyleSheet("color:#64748b;")
        v2.addWidget(self.lbl_anchor)
        b_img = QPushButton("🖼 插入图片到该单元格")
        b_img.clicked.connect(self._insert_image)
        v2.addWidget(b_img)
        self.lbl_imgs = QLabel("本次已插入 0 张图片")
        self.lbl_imgs.setStyleSheet("color:#2563eb;")
        self.lbl_imgs.setWordWrap(True)
        v2.addWidget(self.lbl_imgs)
        v2.addStretch(1)
        split.addWidget(g2, stretch=1)

        lay.addLayout(split, stretch=1)

    # ------------------------------------------------ 数据
    def open_file(self):
        p, _ = QFileDialog.getOpenFileName(self, "打开 Excel 工作簿", "",
                                           "Excel 工作簿 (*.xlsx *.xlsm);;所有文件 (*.*)")
        if p:
            self.open_path(p)

    def open_path(self, p):
        try:
            import openpyxl
            self.wb = openpyxl.load_workbook(p)
        except Exception as e:
            QMessageBox.critical(self, "打开失败",
                                 f"无法打开该文件:\n{e}\n(旧版 .xls 请先另存为 .xlsx)")
            return
        self.path = p
        self.lbl_file.setText(p)
        self.cmb_sheet.blockSignals(True)
        self.cmb_sheet.clear()
        self.cmb_sheet.addItems(self.wb.sheetnames)
        self.cmb_sheet.blockSignals(False)
        self._show_sheet()
        self.status_cb(f"已打开 {p}")

    def _show_sheet(self, *_):
        if self.wb is None:
            return
        self.ws = self.wb[self.cmb_sheet.currentText() or self.wb.sheetnames[0]]
        maxr = min(self.ws.max_row or 1, self.PREVIEW_ROWS)
        maxc = min(self.ws.max_column or 1, self.PREVIEW_COLS)
        self.table.blockSignals(True)
        self.table.clear()
        self.table.setRowCount(maxr)
        self.table.setColumnCount(maxc)
        self.table.setHorizontalHeaderLabels([get_column_letter(c) for c in range(1, maxc + 1)])
        self.table.setVerticalHeaderLabels([str(r) for r in range(1, maxr + 1)])
        for r in range(1, maxr + 1):
            for c in range(1, maxc + 1):
                v = self.ws.cell(row=r, column=c).value
                item = QTableWidgetItem("" if v is None else str(v)[:200])
                self.table.setItem(r - 1, c - 1, item)
        self.table.resizeColumnsToContents()
        self.table.blockSignals(False)

    def _on_cell(self, r, c, *_):
        if self.ws is None:
            return
        self.lbl_pos.setText(f"单元格: {get_column_letter(c + 1)}{r + 1}")
        self.lbl_anchor.setText(f"图片锚点: {get_column_letter(c + 1)}{r + 1}")
        v = self.ws.cell(row=r + 1, column=c + 1).value
        self.txt.setPlainText("" if v is None else str(v))

    # ------------------------------------------------ 操作
    def _require(self):
        if self.wb is None:
            QMessageBox.information(self, "提示", "请先打开一个 Excel 文件")
            return False
        return True

    def _selected_cell(self):
        idx = self.table.currentItem()
        if idx is None or self.ws is None:
            QMessageBox.information(self, "提示", "请先打开文件并选中一个单元格")
            return None
        return idx.row() + 1, idx.column() + 1

    def _apply_cell(self):
        if not self._require():
            return
        cell = self._selected_cell()
        if not cell:
            return
        row, col = cell
        new = self.txt.toPlainText()
        try:
            self.ws.cell(row=row, column=col).value = new
        except Exception as e:
            QMessageBox.critical(self, "错误", f"写入失败:\n{e}")
            return
        item = self.table.item(row - 1, col - 1)
        if item is not None:
            item.setText(new)
        self.status_cb(f"已修改单元格 {get_column_letter(col)}{row}（保存后生效）")

    def _insert_image(self):
        if not self._require():
            return
        cell = self._selected_cell()
        if not cell:
            return
        row, col = cell
        p, _ = QFileDialog.getOpenFileName(self, "选择图片", "",
                                           "图片 (*.png *.jpg *.jpeg *.bmp *.gif *.webp);;所有文件 (*.*)")
        if not p:
            return
        try:
            from openpyxl.drawing.image import Image as XLImage
            from PIL import Image as PILImage
            with PILImage.open(p) as im:
                w, h = im.size
            scale = min(1.0, 480.0 / max(1, w))
            img = XLImage(p)
            img.width = int(w * scale)
            img.height = int(h * scale)
            img.anchor = f"{get_column_letter(col)}{row}"
            self.ws.add_image(img)
            self._img_count += 1
        except Exception as e:
            QMessageBox.critical(self, "错误", f"插入图片失败:\n{e}")
            return
        self.lbl_imgs.setText(f"本次已插入 {self._img_count} 张图片（保存后生效）")
        self.status_cb("图片已加入工作表（保存后生效）")

    def save_as(self):
        if not self._require():
            return
        p, _ = QFileDialog.getSaveFileName(self, "另存 Excel 工作簿",
                                           os.path.basename(self.path), "Excel 工作簿 (*.xlsx)")
        if not p:
            return
        try:
            self.wb.save(p)
        except PermissionError:
            QMessageBox.critical(self, "保存失败", "文件被占用（可能正被 Excel 打开），请先关闭后重试。")
            return
        self.status_cb(f"已保存: {p}")
        QMessageBox.information(self, "保存成功", f"已保存到:\n{p}")
